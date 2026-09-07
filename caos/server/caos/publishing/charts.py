"""Closed frozen chart data and PDF vector marks; no financial calculations."""

from __future__ import annotations

import math
from decimal import Decimal
from typing import Any

from ..contracts import ChartRecipe, DocumentChartSection, is_point_recipe

# Same ordered categorical palette as frontend/components/charts/chartRecipe.ts.
PAPER_COLORS = ("2f54c9", "6741a5", "315b96", "5d5d68", "8a3b8f", "2d6682", "4755a0", "70498d")
CHART_HEIGHT = 240


def prepare_chart(section: dict[str, Any]) -> tuple[ChartRecipe | None, str]:
    """Validate the full grammar/table before converting canonical coordinates.

    Excel stores 15 significant decimal digits. Use that common ceiling for
    both drawings; exact presentation text always remains in the adjacent table.
    """
    if not is_point_recipe(section.get("recipe") or {}):
        return None, "Historical field recipe · authoritative data table"
    checked = DocumentChartSection.model_validate(section)
    recipe = ChartRecipe.model_validate(checked.recipe)
    categories = list(dict.fromkeys(point.x for point in recipe.points))
    seen: set[tuple[str, str]] = set()
    previous: dict[str, int] = {}
    for point in recipe.points:
        for raw in (point.y, point.x_value):
            if raw is None:
                continue
            number = float(raw)
            if (not math.isfinite(number) or (number != 0 and abs(number) < 2.2250738585072014e-308)
                    or Decimal(format(number, ".15g")) != Decimal(raw)):
                return None, "Chart unavailable: coordinate exceeds export range or precision; exact table follows."
        if recipe.kind != "scatter":
            key = (point.x, point.series)
            position = categories.index(point.x)
            if key in seen or position < previous.get(point.series, -1):
                return None, "Chart unavailable: repeated or inconsistently ordered categories; exact table follows."
            seen.add(key)
            previous[point.series] = position
    return recipe, ""


def vector_chart(recipe: ChartRecipe) -> tuple[bytes, list[tuple[float, float, str, float]]]:
    """PDF operators and bounded label boxes in a 504 × 240 point reservation.

    Labels are shaped by the caller's pinned Pango stack. Coordinates are in
    PDF space (origin bottom left); category labels abbreviate, never data.
    """
    categories = list(dict.fromkeys(point.x for point in recipe.points))
    series = list(dict.fromkeys(point.series for point in recipe.points))
    category_index = {category: i for i, category in enumerate(categories)}
    series_index = {name: i for i, name in enumerate(series)}
    values = [float(point.y) for point in recipe.points]
    starts = [0.0] * len(values)
    if recipe.kind == "stacked_bar":
        totals: dict[tuple[str, bool], float] = {}
        # Native XLSX stacks global first-seen series outward from zero for
        # each sign. Accumulate in that order; keep draw/table point order.
        stack_order = sorted(range(len(recipe.points)), key=lambda i: series_index[recipe.points[i].series])
        for i in stack_order:
            point = recipe.points[i]
            key = (point.x, values[i] >= 0)
            starts[i] = totals.get(key, 0.0)
            totals[key] = starts[i] + values[i]
    ends = [start + value for start, value in zip(starts, values)]
    low, high = min(0.0, *starts, *ends), max(0.0, *ends)
    if low == high:
        low, high = -1.0, 1.0
    # Scaling before subtraction avoids overflow in a mixed-sign range.
    scale = max(abs(low), abs(high))
    bottom, top, left, right = 52.0, 218.0, 64.0, 492.0

    def y(value: float) -> float:
        return bottom + ((value / scale - low / scale) / (high / scale - low / scale)) * (top - bottom)

    xs = [float(point.x_value) for point in recipe.points] if recipe.kind == "scatter" else []
    xmin, xmax = (min(xs), max(xs)) if xs else (0.0, 0.0)
    xscale = max(abs(xmin), abs(xmax), 1.0)

    def x(index: int) -> float:
        if xs:
            if xmin == xmax:
                return (left + right) / 2
            return left + 4 + ((xs[index] / xscale - xmin / xscale) / (xmax / xscale - xmin / xscale)) * (right - left - 8)
        return left + (category_index[recipe.points[index].x] + 0.5) * (right - left) / len(categories)

    ops = ["/CAOSChart_" + recipe.kind + " BMC", "q", "0.66 0.64 0.60 RG", "0.5 w"]
    labels: list[tuple[float, float, str, float]] = []
    for value in dict.fromkeys((low, 0.0, high)):
        height = y(value)
        ops.append(f"{left:.4f} {height:.4f} m {right:.4f} {height:.4f} l S")
        labels.append((0, height - 5, format(value, ".6g"), 60))
    if xs:
        labels.extend([(left, 32, format(xmin, ".6g"), 100), (right - 100, 32, format(xmax, ".6g"), 100)])
    else:
        step = max(1, math.ceil(len(categories) / 5))
        for i in range(0, len(categories), step):
            text = categories[i]
            width = min(90, (right - left) / min(len(categories), 5))
            labels.append((left + (i + 0.5) * (right - left) / len(categories) - width / 2, 20,
                           text if len(text) <= 18 else text[:17] + "…", width))
    paths: dict[str, tuple[float, float, int]] = {}
    for i, point in enumerate(recipe.points):
        si = series_index[point.series]
        color = PAPER_COLORS[si % len(PAPER_COLORS)]
        rgb = " ".join(f"{int(color[j:j + 2], 16) / 255:.5f}" for j in (0, 2, 4))
        px, py = x(i), y(ends[i])
        ops.extend([f"{rgb} rg {rgb} RG", "1.5 w"])
        if recipe.kind in {"bar", "stacked_bar"}:
            width = (right - left) / len(categories) * 0.7
            if recipe.kind == "bar":
                width /= len(series)
                px += (si - (len(series) - 1) / 2) * width
            base = y(starts[i])
            if values[i] == 0:
                ops.append(f"{px - width / 2:.4f} {base:.4f} m {px + width / 2:.4f} {base:.4f} l S")
            else:
                ops.append(f"{px - width / 2:.4f} {min(base, py):.4f} {width:.4f} {abs(py - base):.4f} re f")
        else:
            if recipe.kind == "line" and point.series in paths:
                previous_x, previous_y, previous_category = paths[point.series]
                if category_index[point.x] == previous_category + 1:
                    ops.append(f"{previous_x:.4f} {previous_y:.4f} m {px:.4f} {py:.4f} l S")
            ops.append(f"{px - 2:.4f} {py - 2:.4f} 4 4 re f")
            paths[point.series] = (px, py, category_index[point.x])
    ops.extend(["Q", "EMC"])
    return ("\n".join(ops) + "\n").encode("ascii"), labels


def add_xlsx_chart(sheet: Any, recipe: ChartRecipe, *, column: int, safe_cell: Any, title: str | None = None) -> None:
    """Native chart references only canonical numeric cells, never display text."""
    from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.utils import get_column_letter

    series = list(dict.fromkeys(point.series for point in recipe.points))
    categories = list(dict.fromkeys(point.x for point in recipe.points))
    category_index = {category: i for i, category in enumerate(categories)}
    series_index = {name: i for i, name in enumerate(series)}
    scatter = recipe.kind == "scatter"
    sheet.cell(1, column, "Canonical X" if scatter else "Category / period")
    for i, name in enumerate(series, start=1):
        sheet.cell(1, column + i, safe_cell(name))
    if scatter:
        for row, point in enumerate(recipe.points, start=2):
            sheet.cell(row, column, float(point.x_value))
            sheet.cell(row, column + series_index[point.series] + 1, float(point.y))
        end_row = len(recipe.points) + 1
        chart = ScatterChart(scatterStyle="marker")
    else:
        for row, category in enumerate(categories, start=2):
            sheet.cell(row, column, safe_cell(category))
        for point in recipe.points:
            sheet.cell(category_index[point.x] + 2, column + series_index[point.series] + 1, float(point.y))
        end_row = len(categories) + 1
        if recipe.kind == "line":
            chart = LineChart()
            chart.smooth = False
        else:
            chart = BarChart()
            chart.type = "col"
            chart.grouping = "stacked" if recipe.kind == "stacked_bar" else "clustered"
            if recipe.kind == "stacked_bar":
                chart.overlap = 100
            values = [float(point.y) for point in recipe.points]
            # Auto-scaling can start near positive values (e.g. 68/72),
            # exaggerating bars. Keep zero visible without clipping stacks.
            if not any(values):
                chart.y_axis.scaling.min, chart.y_axis.scaling.max = -1, 1
            elif min(values) >= 0:
                chart.y_axis.scaling.min = 0
            elif max(values) <= 0:
                chart.y_axis.scaling.max = 0
    for i, name in enumerate(series, start=1):
        values = Reference(sheet, min_col=column + i, min_row=2, max_row=end_row)
        if scatter:
            x_values = Reference(sheet, min_col=column, min_row=2, max_row=end_row)
            chart.series.append(Series(values, x_values, title=safe_cell(name)))
        else:
            chart.add_data(values)
            chart.series[-1].tx = SeriesLabel(v=safe_cell(name))
        entry = chart.series[-1]
        if isinstance(chart, BarChart):
            entry.invertIfNegative = False
        color = PAPER_COLORS[(i - 1) % len(PAPER_COLORS)]
        entry.graphicalProperties.solidFill = color
        entry.graphicalProperties.line.solidFill = color
        if scatter or recipe.kind == "line":
            entry.smooth = False
            entry.marker.symbol = "square"
            entry.marker.size = 5
            entry.marker.graphicalProperties.solidFill = color
            entry.marker.graphicalProperties.line.solidFill = color
            if scatter:
                entry.graphicalProperties.line.solidFill = None
                entry.graphicalProperties.line.noFill = True
    if not scatter:
        chart.set_categories(Reference(sheet, min_col=column, min_row=2, max_row=end_row))
    chart.title = safe_cell(title if title is not None else sheet.title)
    chart.x_axis.crosses = "autoZero"
    chart.x_axis.tickLblPos = "low"
    chart.y_axis.crosses = "autoZero"
    chart.y_axis.title = safe_cell(recipe.unit)
    chart.height, chart.width = 10, 20
    chart.display_blanks = "gap"
    chart.legend.position = "b"
    sheet.add_chart(chart, f"{get_column_letter(column)}{end_row + 3}")

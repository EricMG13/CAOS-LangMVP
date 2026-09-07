"""Frozen chart rendering: canonical coordinates, exact tables and old bytes."""

import copy
import io
import json
import os
from pathlib import Path
from zipfile import ZipFile

import pytest

from caos.contracts import ChartRecipe
from caos.publishing.renderers import render_frozen_export


def chart_section(kind="bar"):
    recipe = {
        "schema_version": "caos.chart.v1", "recipe_id": f"fixture-{kind}",
        "kind": kind, "unit": "USD m", "points": [
            {"x": category, "y": value, "display": f"{value} million", "series": series,
             "source_ids": ["src-reviewed"], **({"x_value": str(index + 1)} if kind == "scatter" else {})}
            for index, (category, series, value) in enumerate([
                ("FY2025 long reviewed category label", "Base", "12.5"),
                ("FY2026", "Base", "-3"), ("FY2027", "Base", "0"),
                ("FY2025 long reviewed category label", "Downside", "8"),
                ("FY2026", "Downside", "-5"), ("FY2027", "Downside", "2"),
            ])
        ],
    }
    columns, rows = ChartRecipe.model_validate(recipe).accessible_table()
    return {"kind": "chart", "section_id": f"chart-{kind}", "title": f"Reviewed {kind}",
            "page": "Financials", "editable": False,
            "origin": {"kind": "ARTIFACT", "authority_id": "artifact-reviewed", "block_ids": ["block-reviewed"]},
            "recipe": recipe, "accessible_columns": columns, "accessible_rows": rows}


def chart_payload(kinds=("line", "bar", "stacked_bar", "scatter"), version="caos.deliverable-renderer.v4"):
    return {"renderer": {"version": version}, "template": {"template_version": "caos.deliverable-template.v2"}, "preview_digest": "reviewed-digest",
            "publication": {"masthead": {"issuer": "Chart Fixture", "report_type": "Credit Report",
                "renderer_version": version, "model_identity": "MODEL_BUILD model-reviewed",
                "approval_state": "PENDING APPROVAL"},
                "pages": [{"name": "Financials", "sections": [chart_section(kind) for kind in kinds]}]}}


@pytest.mark.parametrize("kind", ["line", "bar", "stacked_bar", "scatter"])
def test_native_chart_uses_canonical_coordinates_and_retains_exact_table(kind):
    from openpyxl import load_workbook

    payload = chart_payload((kind,))
    before = copy.deepcopy(payload)
    data = render_frozen_export(payload, "xlsx")
    book = load_workbook(io.BytesIO(data))
    sheet = book[f"Reviewed {kind}"]
    assert len(sheet._charts) == 1
    if kind in {"line", "scatter"}:
        assert all(series.smooth is False for series in sheet._charts[0].series)
    chart = sheet._charts[0]
    assert chart.x_axis.crosses == chart.y_axis.crosses == "autoZero"
    assert chart.x_axis.tickLblPos == "low"
    if kind in {"bar", "stacked_bar"}:
        assert all(series.invertIfNegative is False for series in chart.series)
    if kind == "scatter":
        assert chart.scatterStyle == "marker"
        assert all(series.graphicalProperties.line.noFill and series.graphicalProperties.line.solidFill is None for series in chart.series)
    section = payload["publication"]["pages"][0]["sections"][0]
    for i, row in enumerate([section["accessible_columns"], *section["accessible_rows"]], start=1):
        assert [sheet.cell(i, j).value for j in range(1, len(row) + 1)] == row
    assert payload == before
    assert render_frozen_export(payload, "xlsx") == data


@pytest.mark.parametrize("kind", ["bar", "stacked_bar"])
@pytest.mark.parametrize("values,bounds", [(["68", "72"], (0, None)), (["-68", "-72"], (None, 0)), (["0", "0"], (-1, 1))])
def test_native_bar_zero_baseline_and_full_chart_title(kind, values, bounds):
    from openpyxl import load_workbook

    payload = chart_payload((kind,))
    section = payload["publication"]["pages"][0]["sections"][0]
    section["title"] = "Debt maturities — carrying value, FY2024_Q4"
    section["recipe"]["points"] = section["recipe"]["points"][:2]
    for point, value in zip(section["recipe"]["points"], values):
        point.update(y=value, display=value)
    section["accessible_columns"], section["accessible_rows"] = ChartRecipe.model_validate(section["recipe"]).accessible_table()
    book = load_workbook(io.BytesIO(render_frozen_export(payload, "xlsx")))
    sheet = next(sheet for sheet in book if sheet._charts)
    chart = sheet._charts[0]
    assert (chart.y_axis.scaling.min, chart.y_axis.scaling.max) == bounds
    assert len(sheet.title) <= 31
    assert chart.title.tx.rich.p[0].r[0].t == section["title"]


@pytest.mark.parametrize("value", ["1234567890123456", "0.1234567890123456", "9" * 128])
def test_unrepresentable_coordinates_keep_the_exact_table(value):
    from openpyxl import load_workbook

    payload = chart_payload(("bar",))
    section = payload["publication"]["pages"][0]["sections"][0]
    section["recipe"]["points"][0].update(y=value, display=value)
    section["accessible_columns"], section["accessible_rows"] = ChartRecipe.model_validate(section["recipe"]).accessible_table()
    book = load_workbook(io.BytesIO(render_frozen_export(payload, "xlsx")))
    sheet = book["Reviewed bar"]
    assert not sheet._charts
    assert sheet["C2"].value == value and sheet["C2"].data_type == "s"
    assert "Chart unavailable" in " ".join(str(c.value) for row in sheet for c in row)


def test_pdf_draws_four_closed_kinds_and_markdown_names_text_representation():
    from pypdf import PdfReader

    payload = chart_payload()
    pages = PdfReader(io.BytesIO(render_frozen_export(payload, "pdf"))).pages
    assert sum(page.get_contents().get_data().count(b"/CAOSChart_") for page in pages) == 4
    text = " ".join(page.extract_text() for page in pages)
    assert "src-reviewed" in text and "artifact-reviewed" in text and "model-reviewed" in text
    markdown = render_frozen_export(payload, "md").decode()
    assert "text representation" in markdown
    assert "stacked_bar" in markdown


def test_old_renderer_and_schema_labelled_field_recipe_stay_table_only():
    from openpyxl import load_workbook

    for version in (None, "caos.deliverable-renderer.v3"):
        payload = chart_payload(("bar",), version)
        if version is None:
            payload.pop("renderer")
        book = load_workbook(io.BytesIO(render_frozen_export(payload, "xlsx")))
        assert not book["Reviewed bar"]._charts
    payload = chart_payload(("bar",))
    payload["publication"]["pages"][0]["sections"][0]["recipe"] = {
        "schema_version": "caos.chart.v1", "kind": "bar", "fields": ["metric"]}
    book = load_workbook(io.BytesIO(render_frozen_export(payload, "xlsx")))
    assert not book["Reviewed bar"]._charts


def test_reviewed_chart_export_goldens(tmp_path):
    from openpyxl import load_workbook
    from pypdf import PdfReader

    payload = chart_payload()
    output = Path(os.environ.get("CAOS_CHART_INSPECTION_DIR", tmp_path))
    output.mkdir(parents=True, exist_ok=True)
    exports = {fmt: render_frozen_export(payload, fmt) for fmt in ("md", "pdf", "xlsx")}
    for fmt, content in exports.items():
        (output / f"reviewed-charts.{fmt}").write_bytes(content)
    (output / "reviewed-payload.json").write_text(json.dumps(payload, indent=2))
    book = load_workbook(io.BytesIO(exports["xlsx"]))
    dump = {sheet.title: {"rows": list(sheet.values), "charts": [str(chart.to_tree().tag) for chart in sheet._charts]}
            for sheet in book}
    pages = PdfReader(io.BytesIO(exports["pdf"])).pages
    observed = {
        "md": exports["md"].decode(),
        "xlsx": json.dumps({"sheets": dump, "chart_xml": {
            name: ZipFile(io.BytesIO(exports["xlsx"])).read(name).decode()
            for name in ZipFile(io.BytesIO(exports["xlsx"])).namelist() if name.startswith("xl/charts/chart")}}, indent=1),
        "pdf": json.dumps({"renderer": payload["renderer"], "pages": [{"text": page.extract_text(),
            "charts": page.get_contents().get_data().count(b"/CAOSChart_")} for page in pages]}, indent=1),
    }
    (output / "sheet-inventory.json").write_text(json.dumps(dump, indent=2))
    golden_dir = Path(__file__).resolve().parents[1] / "fixtures/deliverables/publication_charts"
    for fmt, value in observed.items():
        target = golden_dir / f"reviewed.{fmt}.golden"
        if os.environ.get("CAOS_REGENERATE_GOLDENS") == "1":
            golden_dir.mkdir(parents=True, exist_ok=True)
            target.write_text(value)
        assert target.exists(), "Inspect every chart page and sheet before deliberate golden regeneration"
        assert target.read_text() == value


def test_formula_like_labels_are_text_and_series_reference_only_coordinate_cells():
    from openpyxl import load_workbook

    payload = chart_payload(("scatter",))
    section = payload["publication"]["pages"][0]["sections"][0]
    section["title"] = "=Reviewed"
    section["recipe"]["points"][0].update(x="=1+1", series="@Base", display="+999", y="12.5", x_value="-2")
    section["accessible_columns"], section["accessible_rows"] = ChartRecipe.model_validate(section["recipe"]).accessible_table()
    book = load_workbook(io.BytesIO(render_frozen_export(payload, "xlsx")))
    sheet = book["=Reviewed"]
    assert sheet["A2"].value == "=1+1" and sheet["A2"].data_type == "s"
    assert sheet["C2"].value == "+999" and sheet["C2"].data_type == "s"
    assert sheet["H2"].value == -2 and sheet["I2"].value == 12.5
    chart, = sheet._charts
    assert chart.series[0].xVal.numRef.f.endswith("!$H$2:$H$7")
    assert chart.series[0].yVal.numRef.f.endswith("!$I$2:$I$7")
    assert not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row)


def test_sparse_line_series_preserve_gaps_in_pdf_and_xlsx():
    from caos.publishing.charts import prepare_chart, vector_chart
    from openpyxl import load_workbook

    payload = chart_payload(("line",))
    section = payload["publication"]["pages"][0]["sections"][0]
    section["recipe"]["points"] = [section["recipe"]["points"][i] for i in (0, 4, 2)]
    section["accessible_columns"], section["accessible_rows"] = ChartRecipe.model_validate(section["recipe"]).accessible_table()
    recipe, reason = prepare_chart(section)
    assert not reason
    operators, _ = vector_chart(recipe)
    # Only the three axis lines; no invented segment across the missing value.
    assert operators.count(b" l S") == 3
    sheet = load_workbook(io.BytesIO(render_frozen_export(payload, "xlsx")))["Reviewed line"]
    assert [sheet.cell(row, 8).value for row in range(2, 5)] == [12.5, None, 0]
    assert sheet._charts[0].display_blanks == "gap"


@pytest.mark.parametrize("rows", [
    [("Q1", "A", "1"), ("Q1", "B", "2"), ("Q2", "B", "2"), ("Q2", "A", "1")],
    [("Q1", "A", "-1"), ("Q1", "B", "-2"), ("Q2", "B", "-2"), ("Q2", "A", "-1")],
    [("Q1", "A", "1"), ("Q1", "B", "-2"), ("Q1", "C", "0"),
     ("Q2", "C", "-3"), ("Q2", "B", "2"), ("Q2", "A", "-1"),
     ("Q3", "B", "0"), ("Q3", "C", "-4"), ("Q3", "A", "-2")],
])
def test_signed_stacks_follow_global_series_order_without_reordering_exact_rows(rows):
    import re
    from caos.publishing.charts import prepare_chart, vector_chart
    from openpyxl import load_workbook

    payload = chart_payload(("stacked_bar",))
    section = payload["publication"]["pages"][0]["sections"][0]
    section["recipe"]["points"] = [{"x": category, "series": series, "y": value,
        "display": value, "source_ids": ["src-reviewed"]} for category, series, value in rows]
    section["accessible_columns"], section["accessible_rows"] = ChartRecipe.model_validate(section["recipe"]).accessible_table()
    before = copy.deepcopy(payload)
    recipe, reason = prepare_chart(section)
    assert not reason
    series = list(dict.fromkeys(row[1] for row in rows))
    categories = list(dict.fromkeys(row[0] for row in rows))
    intervals = {}
    for category in categories:
        positive = negative = 0.0
        for name in series:
            value = next((float(y) for x, s, y in rows if (x, s) == (category, name)), None)
            if value is None:
                continue
            start = negative if value < 0 else positive
            intervals[category, name] = (start, start + value)
            if value < 0:
                negative += value
            else:
                positive += value
    low = min(0, *(min(pair) for pair in intervals.values()))
    high = max(0, *(max(pair) for pair in intervals.values()))
    operators, _ = vector_chart(recipe)
    rectangles = re.findall(rb"([-\d.]+) ([-\d.]+) ([-\d.]+) ([-\d.]+) re f", operators)
    expected = [intervals[x, s] for x, s, value in rows if float(value) != 0]
    assert len(rectangles) == len(expected)
    for rectangle, (start, end) in zip(rectangles, expected):
        assert float(rectangle[1]) == pytest.approx(52 + (min(start, end) - low) / (high - low) * 166, abs=0.0001)
        assert float(rectangle[3]) == pytest.approx(abs(end - start) / (high - low) * 166, abs=0.0001)
    sheet = load_workbook(io.BytesIO(render_frozen_export(payload, "xlsx")))["Reviewed stacked_bar"]
    chart = sheet._charts[0]
    assert chart.grouping == "stacked" and chart.overlap == 100
    assert [entry.tx.v for entry in chart.series] == series
    assert all(entry.invertIfNegative is False for entry in chart.series)
    for i, category in enumerate(categories, 2):
        assert sheet.cell(i, 7).value == category
        for j, name in enumerate(series, 8):
            assert sheet.cell(i, j).value == next((float(y) for x, s, y in rows if (x, s) == (category, name)), None)
    assert payload == before
    assert [[sheet.cell(i, j).value for j in range(1, 6)] for i in range(2, len(rows) + 2)] == section["accessible_rows"]
    if os.environ.get("CAOS_CHART_INSPECTION_DIR"):
        directory = Path(os.environ["CAOS_CHART_INSPECTION_DIR"])
        directory.mkdir(parents=True, exist_ok=True)
        name = "positive" if all(float(y) >= 0 for _, _, y in rows) else "negative" if all(float(y) <= 0 for _, _, y in rows) else "mixed"
        for fmt in ("pdf", "xlsx"):
            (directory / f"stack-{name}.{fmt}").write_bytes(render_frozen_export(payload, fmt))


@pytest.mark.parametrize("kind", ["line", "bar", "stacked_bar", "scatter"])
def test_zero_single_point_and_numeric_extremes_produce_bounded_marks(kind):
    from caos.publishing.charts import prepare_chart, vector_chart

    section = chart_section(kind)
    section["recipe"]["points"] = [section["recipe"]["points"][0]]
    for value in ("0", "-1", "1" + "0" * 120, "0." + "0" * 120 + "1"):
        section["recipe"]["points"][0].update(y=value, display=value)
        section["accessible_columns"], section["accessible_rows"] = ChartRecipe.model_validate(section["recipe"]).accessible_table()
        recipe, reason = prepare_chart(section)
        assert not reason
        operators, labels = vector_chart(recipe)
        assert b"nan" not in operators and b"inf" not in operators
        assert all(0 <= x <= 504 and 0 <= y <= 240 and 0 < width <= 504 - x for x, y, _, width in labels)


@pytest.mark.parametrize("change", ["table", "unknown", "coordinate"])
def test_closed_recipe_and_equivalent_table_are_revalidated_at_export(change):
    from caos.publishing.charts import prepare_chart

    section = chart_section()
    if change == "table":
        section["accessible_rows"][0][2] = "forged"
    elif change == "unknown":
        section["recipe"]["callback"] = "arbitrary"
    else:
        section["recipe"]["points"][0]["y"] = "NaN"
    with pytest.raises(ValueError):
        prepare_chart(section)


def test_new_v1_freezes_select_v4_but_queued_historical_exports_rebuild_as_v3(tmp_path, store, monkeypatch):
    from test_deliverables_spec import make_service, seed_ready_case, draft_request, sign_min, freeze_request

    service = make_service(store, tmp_path / "vault")
    case, source, _ = seed_ready_case(service, store)
    template = service.templates(template_version="caos.deliverable-template.v1")["RELATIVE_VALUE"]
    revision = service.save_draft(case["id"], "RELATIVE_VALUE", draft_request(template, source), actor="analyst")
    sign_min(service, case["id"], revision)
    with monkeypatch.context() as historical:
        historical.setattr("caos.deliverables.service.RENDERER_VERSION", "caos.deliverable-renderer.v3")
        job = service.freeze(case["id"], freeze_request(revision), actor="analyst")
    assert job["status"] == "QUEUED"
    assert service.run_pending_freezes() == 1
    frozen = service.frozen_record_for_job(case["id"], job["job_id"])
    assert frozen["payload"]["renderer"]["version"] == "caos.deliverable-renderer.v3"
    for fmt in ("md", "pdf", "xlsx"):
        assert service.export(frozen["deliverable_id"], fmt)[0] == render_frozen_export(frozen["payload"], fmt)
    next_revision = service.save_draft(case["id"], "RELATIVE_VALUE", draft_request(template, source, expected_version=1), actor="analyst")
    sign_min(service, case["id"], next_revision)
    current = service.freeze(case["id"], freeze_request(next_revision), actor="analyst")
    service.run_pending_freezes()
    assert service.frozen_record_for_job(case["id"], current["job_id"])["payload"]["renderer"]["version"] == "caos.deliverable-renderer.v4"


def test_v2_reviewed_sections_freeze_file_and_retain_all_bytes_after_withdrawal(tmp_path, store, monkeypatch):
    from test_deliverables_spec import (make_service, seed_ready_case, report_artifacts, draft_request, freeze_now,
                                        seed_model, revision_selection, add_approver, file_request)

    service = make_service(store, tmp_path / "vault")
    case, source, _ = seed_ready_case(service, store)
    monkeypatch.setattr(service, "_accepted_artifacts", lambda _: report_artifacts(source["id"]))
    model = seed_model(service, case)
    template = service.templates()["FULL_CREDIT"]
    revision = service.save_draft(case["id"], "FULL_CREDIT", draft_request(template, source,
        model_selection=revision_selection(model)), actor="analyst")
    reviewed = copy.deepcopy(revision["content"])
    frozen = freeze_now(service, case["id"], revision)
    payload = frozen["payload"]
    assert payload["content"] == reviewed
    sections = {section["section_id"]: section for page in payload["publication"]["pages"] for section in page["sections"]}
    assert all(sections[section["section_id"]] == section for section in reviewed["document_sections"])
    assert payload["content"]["model_identity"] == reviewed["model_identity"]
    assert {(e["source_id"], tuple(e["block_ids"])) for e in payload["evidence"]} == {
        (e["source_id"], tuple(e["block_ids"])) for e in reviewed["citation_union"]}
    before = {fmt: service.export(frozen["deliverable_id"], fmt) for fmt in ("md", "pdf", "xlsx")}
    if os.environ.get("CAOS_CHART_INSPECTION_DIR"):
        directory = Path(os.environ["CAOS_CHART_INSPECTION_DIR"])
        directory.mkdir(parents=True, exist_ok=True)
        for fmt, (content, _) in before.items():
            (directory / f"v2-frozen.{fmt}").write_bytes(content)
        (directory / "v2-frozen-payload.json").write_text(json.dumps(payload, indent=2))
    from pypdf import PdfReader
    pdf_text = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(before["pdf"][0])).pages)
    expected_keys = [f"Series key · ■ {series}" for section in sections.values() if section["kind"] == "chart"
                     for series in dict.fromkeys(point["series"] for point in section["recipe"].get("points", []))]
    assert all(key in pdf_text for key in expected_keys), [key for key in expected_keys if key not in pdf_text]
    add_approver(store, case)
    filed = service.approve_filing(case["id"], frozen["deliverable_id"], file_request(frozen), actor="approver-user")
    monkeypatch.setattr("caos.deliverables.service.RENDERER_VERSION", "future-update")
    store.withdraw(case["id"], source["id"], "analyst")
    assert {fmt: service.export(filed["deliverable_id"], fmt) for fmt in before} == before
    assert service.frozen_record(case["id"], filed["deliverable_id"])["payload"] == payload


def test_control_register_paginates_every_source_and_evidence_row():
    from caos.publishing.document import build_publication
    from test_deliverables_spec import OPINION

    sources = [{"id": f"source-{i}", "filename": f"document-{i}", "sha256": f"{i:064x}"} for i in range(501)]
    evidence = [{"source_id": source["id"], "block_ids": ["block"], "withdrawn": False} for source in sources]
    payload = {"content": {"document_sections": [], "blocks": []}, "authority": {}, "case_id": "case-fixture",
               "evidence": evidence, "template": {"title": "Credit Report"}, "pathway": "FULL_CREDIT",
               "draft": {"version": 1, "digest": "digest"}, "input_fingerprint": "fingerprint",
               "methodology": {"build_id": "methodology"}, "renderer": {"version": "caos.deliverable-renderer.v4"}}
    publication = build_publication(payload=payload, opinion={**OPINION, "opinion_id": "opinion-fixture",
        "signed_by": "analyst", "signed_at": "2026-09-07", "opinion_digest": "digest"}, case={"issuer": "Fixture"},
        deliverable_id="deliverable-fixture", provider_identity=None, accepted_at=None, run_id=None,
        sources=sources, dispositions={})
    for prefix in ("source_document_register", "registered_evidence_inventory"):
        sections = [s for page in publication["pages"] for s in page["sections"] if s["section_id"].startswith(prefix)]
        assert [len(s["rows"]) for s in sections] == [500, 1]
        assert [row[0] for s in sections for row in s["rows"]] == [source["id"] for source in sources]
    assert payload["evidence"] == evidence


def test_chart_reservations_remain_whole_and_headers_repeat_only_after_page_breaks():
    from pypdf import PdfReader

    payload = chart_payload()
    pages = PdfReader(io.BytesIO(render_frozen_export(payload, "pdf"))).pages
    assert len(pages) == 4
    for page in pages:
        content = page.get_contents().get_data()
        assert content.count(b"/CAOSChart_") == 1
        text = page.extract_text()
        assert text.count("Category / period") == 1
        assert "PAGE " in text and "src-reviewed" in text
